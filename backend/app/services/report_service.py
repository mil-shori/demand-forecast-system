"""
月次経理レポート生成サービス（Excel / CSV）
"""
import calendar
import io
import logging
from datetime import date
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models import Expense, JournalEntry, JournalEntryStatus
from app.services.accounting_service import aggregate_sales

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
HEADER_FONT = Font(bold=True)


def _month_range(year: int, month: int) -> tuple[date, date]:
    last_day = calendar.monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last_day)


def _write_sheet(ws, headers: List[str], rows: List[List[Any]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append(row)
    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(len(str(header)) * 2, 14)


def collect_monthly_data(db: Session, year: int, month: int) -> Dict[str, Any]:
    """月次レポート用のデータを収集する"""
    start, end = _month_range(year, month)

    sales = aggregate_sales(db, start, end, granularity="daily")

    expenses = (
        db.query(Expense)
        .filter(Expense.expense_date >= start, Expense.expense_date <= end)
        .order_by(Expense.expense_date)
        .all()
    )

    expense_by_category = (
        db.query(
            Expense.category,
            func.sum(Expense.amount).label("amount"),
            func.count(Expense.id),
        )
        .filter(Expense.expense_date >= start, Expense.expense_date <= end)
        .group_by(Expense.category)
        .all()
    )

    journal_entries = (
        db.query(JournalEntry)
        .filter(JournalEntry.entry_date >= start, JournalEntry.entry_date <= end)
        .order_by(JournalEntry.entry_date)
        .all()
    )

    return {
        "start": start,
        "end": end,
        "sales": sales,
        "expenses": expenses,
        "expense_by_category": expense_by_category,
        "journal_entries": journal_entries,
        "total_sales": sum(s["amount"] for s in sales),
        "total_expenses": sum(e.amount for e in expenses),
    }


def generate_monthly_report_xlsx(
    db: Session,
    year: int,
    month: int,
    trial_pl: Optional[Dict[str, Any]] = None,
) -> io.BytesIO:
    """月次レポートをExcel（4シート）で生成する"""
    data = collect_monthly_data(db, year, month)
    wb = Workbook()

    # シート1: 売上集計
    ws = wb.active
    ws.title = "売上集計"
    sales_rows = [
        [s["period"], s["order_type"], float(s["amount"]), s["order_count"]]
        for s in data["sales"]
    ]
    sales_rows.append(
        [
            "合計",
            "",
            float(data["total_sales"]),
            sum(s["order_count"] for s in data["sales"]),
        ]
    )
    _write_sheet(ws, ["日付", "注文タイプ", "売上金額", "注文件数"], sales_rows)

    # シート2: 経費一覧
    ws = wb.create_sheet("経費一覧")
    expense_rows = [
        [
            e.expense_date.isoformat(),
            e.category,
            float(e.amount),
            e.payment_method,
            e.partner_name or "",
            e.description or "",
            e.status.value,
        ]
        for e in data["expenses"]
    ]
    expense_rows.append(["合計", "", float(data["total_expenses"]), "", "", "", ""])
    for category, amount, count in data["expense_by_category"]:
        expense_rows.append(
            [f"小計: {category}", "", float(amount), f"{count}件", "", "", ""]
        )
    _write_sheet(
        ws,
        ["日付", "カテゴリ", "金額", "支払方法", "取引先", "摘要", "ステータス"],
        expense_rows,
    )

    # シート3: freee試算表（接続時のみ）
    ws = wb.create_sheet("freee試算表")
    if trial_pl and trial_pl.get("balances"):
        pl_rows = [
            [
                b.get("account_item_name") or b.get("account_category_name") or "",
                b.get("closing_balance", 0),
            ]
            for b in trial_pl["balances"]
        ]
        _write_sheet(ws, ["勘定科目", "残高"], pl_rows)
    else:
        ws.append(["freee未接続のため試算表は出力されていません。"])

    # シート4: 同期状況
    ws = wb.create_sheet("同期状況")
    synced = sum(
        1 for j in data["journal_entries"] if j.status == JournalEntryStatus.SYNCED
    )
    failed = sum(
        1 for j in data["journal_entries"] if j.status == JournalEntryStatus.FAILED
    )
    sync_rows = [
        [
            j.entry_date.isoformat(),
            j.entry_type.value,
            j.description or "",
            float(j.amount),
            j.status.value,
            j.freee_deal_id or "",
            j.error_message or "",
        ]
        for j in data["journal_entries"]
    ]
    sync_rows.append([f"SYNCED: {synced}件 / FAILED: {failed}件", "", "", "", "", "", ""])
    _write_sheet(
        ws,
        ["日付", "種別", "摘要", "金額", "ステータス", "freee取引ID", "エラー"],
        sync_rows,
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_monthly_report_csv(db: Session, year: int, month: int) -> io.BytesIO:
    """月次の売上集計CSVを生成する（Excelで開けるようBOM付きUTF-8）"""
    data = collect_monthly_data(db, year, month)
    df = pd.DataFrame(
        [
            {
                "日付": s["period"],
                "注文タイプ": s["order_type"],
                "売上金額": float(s["amount"]),
                "注文件数": s["order_count"],
            }
            for s in data["sales"]
        ]
    )
    buffer = io.BytesIO()
    buffer.write(df.to_csv(index=False).encode("utf-8-sig"))
    buffer.seek(0)
    return buffer
