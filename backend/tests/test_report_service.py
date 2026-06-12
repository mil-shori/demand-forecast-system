"""
report_service のテスト（xlsx読み戻し検証・CSV出力）
"""
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.models import Expense, ExpenseStatus, Order, OrderType
from app.services import report_service


def _seed(db):
    db.add_all(
        [
            Order(
                order_id="O1",
                user_id="U1",
                order_date=date(2026, 5, 1),
                order_type=OrderType.SUBSCRIPTION,
                total_amount=Decimal("5000"),
            ),
            Order(
                order_id="O2",
                user_id="U2",
                order_date=date(2026, 5, 2),
                order_type=OrderType.ONEOFF,
                total_amount=Decimal("1500"),
            ),
            # 集計対象外（6月）
            Order(
                order_id="O3",
                user_id="U3",
                order_date=date(2026, 6, 1),
                order_type=OrderType.ONEOFF,
                total_amount=Decimal("9999"),
            ),
            Expense(
                expense_date=date(2026, 5, 10),
                amount=Decimal("1200"),
                category="交通費",
                description="タクシー代",
                payment_method="cash",
                status=ExpenseStatus.DRAFT,
            ),
            Expense(
                expense_date=date(2026, 5, 15),
                amount=Decimal("800"),
                category="会議費",
                payment_method="credit_card",
                status=ExpenseStatus.SYNCED,
                freee_deal_id=1,
            ),
        ]
    )
    db.commit()


def test_collect_monthly_data(db_session):
    _seed(db_session)
    data = report_service.collect_monthly_data(db_session, 2026, 5)
    assert data["total_sales"] == Decimal("6500.00")
    assert data["total_expenses"] == Decimal("2000.00")
    assert len(data["expenses"]) == 2


def test_generate_xlsx_sheets_and_values(db_session):
    _seed(db_session)
    buffer = report_service.generate_monthly_report_xlsx(db_session, 2026, 5)
    wb = load_workbook(buffer)

    assert wb.sheetnames == ["売上集計", "経費一覧", "freee試算表", "同期状況"]

    sales = wb["売上集計"]
    rows = list(sales.iter_rows(values_only=True))
    assert rows[0] == ("日付", "注文タイプ", "売上金額", "注文件数")
    # 合計行（最終行）
    assert rows[-1][0] == "合計"
    assert rows[-1][2] == 6500.0
    # 6月分は含まれない
    assert all("2026-06" not in str(r[0]) for r in rows[1:])

    expenses = wb["経費一覧"]
    expense_rows = list(expenses.iter_rows(values_only=True))
    categories = [r[1] for r in expense_rows[1:]]
    assert "交通費" in categories and "会議費" in categories

    # freee未接続メッセージ
    trial = wb["freee試算表"]
    assert "freee未接続" in trial["A1"].value


def test_generate_xlsx_with_trial_pl(db_session):
    _seed(db_session)
    trial_pl = {
        "balances": [
            {"account_item_name": "売上高", "closing_balance": 6500},
            {"account_item_name": "旅費交通費", "closing_balance": 1200},
        ]
    }
    buffer = report_service.generate_monthly_report_xlsx(db_session, 2026, 5, trial_pl)
    wb = load_workbook(buffer)
    trial = wb["freee試算表"]
    rows = list(trial.iter_rows(values_only=True))
    assert rows[0] == ("勘定科目", "残高")
    assert ("売上高", 6500) in rows


def test_generate_csv(db_session):
    _seed(db_session)
    buffer = report_service.generate_monthly_report_csv(db_session, 2026, 5)
    content = buffer.read().decode("utf-8-sig")
    assert "売上金額" in content
    assert "5000.0" in content
    assert "9999" not in content  # 6月分は含まれない
